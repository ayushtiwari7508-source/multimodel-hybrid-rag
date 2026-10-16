"""Document loaders for various file formats."""
from __future__ import annotations

import io
import re
from collections.abc import Iterable
from pathlib import Path

import pdfplumber
from PIL import Image
from pytesseract import image_to_string

from backend.app.core.config import get_settings
from backend.app.core.observability import get_logger
from backend.app.generation.llm import OllamaClient
from backend.app.models.domain import DocElement, LoadedDocument

log = get_logger(__name__)


def ocr_image(img: Image.Image) -> str:
    try:
        return image_to_string(img).strip()
    except Exception as e:
        log.warning(f"OCR failed: {e}")
        return ""


def describe_image_ollama(img: Image.Image, vision_model: str, ollama: OllamaClient) -> str:
    try:
        buf = io.BytesIO()
        if img.mode in ("RGBA", "LA", "P"):
            img = img.convert("RGB")
        img.save(buf, format="PNG")
        out = ollama.generate(
            prompt=(
                "Describe this image precisely in detail. Include any text, labels, values, "
                "axis labels, and key visual elements you see. If it is a chart, "
                "describe axes, series, and trends."
            ),
            model=vision_model,
            images=[buf.getvalue()],
            options={"temperature": 0.1, "num_predict": 400},
        )
        return out.strip() if isinstance(out, str) else ""
    except Exception as e:
        log.warning(f"Vision describe failed: {e}")
        return ""


def extract_tables_pdfplumber(page) -> list[str]:
    tables: list[str] = []
    try:
        for tbl in page.extract_tables() or []:
            rows = []
            for r in tbl:
                rows.append(" | ".join((c or "").strip() for c in r))
            if rows:
                tables.append("\n".join(rows))
    except Exception:
        pass
    return tables


def extract_tables_camelot(pdf_path: str) -> list[tuple[int, str]]:
    try:
        import camelot

        tables = camelot.read_pdf(pdf_path, pages="all", flavor="lattice")
        out: list[tuple[int, str]] = []
        for t in enumerate(tables):
            out.append((t[1].page, t[1].df.to_csv(index=False)))
        return out
    except Exception:
        return []


def extract_equations_heuristic(text: str) -> tuple[str, list[str]]:
    eqs: list[str] = []
    cleaned_lines: list[str] = []
    eq_patterns = [
        r"\$\$(.+?)\$\$",
        r"\$(.+?)\$",
        r"\\\[(.+?)\\\]",
        r"\\\((.+?)\\\)",
    ]
    for line in text.split("\n"):
        matched = False
        for pat in eq_patterns:
            for m in re.findall(pat, line):
                eqs.append(m.strip())
                matched = True
        if not matched:
            cleaned_lines.append(line)
    return "\n".join(cleaned_lines), eqs


def load_pdf(path: str, ollama: OllamaClient) -> LoadedDocument:
    cfg = get_settings()
    doc_id = Path(path).stem
    doc = LoadedDocument(source=path, doc_id=doc_id)

    camelot_tables = {}
    try:
        for page_no, csv in extract_tables_camelot(path):
            camelot_tables.setdefault(page_no, []).append(csv)
    except Exception:
        pass

    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            cleaned, eqs = extract_equations_heuristic(text)
            has_text = bool(cleaned.strip())
            if has_text:
                doc.elements.append(DocElement(type="text", content=cleaned, page=i, section=f"page {i}"))
            for eq in eqs:
                doc.elements.append(DocElement(type="equation", content=eq, page=i, section=f"page {i}"))
            for tbl in extract_tables_pdfplumber(page):
                doc.elements.append(DocElement(type="table", content=tbl, page=i, section=f"page {i}"))
            for csv in camelot_tables.get(i, []):
                doc.elements.append(DocElement(type="table", content=csv, page=i, section=f"page {i}"))

            if not cfg.enable_vision:
                continue
            if cfg.vision_only_when_no_text and has_text:
                continue
            try:
                for img in page.images:
                    try:
                        x0, y0, x1, y1 = img["x0"], img["top"], img["x1"], img["bottom"]
                        crop = page.crop((x0, y0, x1, y1)).to_image(resolution=200)
                        pil_img = crop.original
                    except Exception:
                        continue
                    desc = describe_image_ollama(pil_img, cfg.vision_model, ollama)
                    if desc:
                        doc.elements.append(DocElement(type="image", content=desc, page=i, section=f"page {i}"))
                    ocr = ocr_image(pil_img)
                    if ocr and len(ocr) > 3:
                        doc.elements.append(DocElement(type="ocr", content=ocr, page=i, section=f"page {i}"))
            except Exception as e:
                log.debug(f"image extraction issue p{i}: {e}")
    return doc


def load_markdown(path: str) -> LoadedDocument:
    text = Path(path).read_text(encoding="utf-8", errors="ignore")
    doc = LoadedDocument(source=path, doc_id=Path(path).stem)
    current_section = "root"
    for line in text.split("\n"):
        if line.startswith("#"):
            current_section = line.lstrip("#").strip() or current_section
        if line.strip():
            doc.elements.append(DocElement(type="text", content=line, section=current_section))
    return doc


def load_text(path: str) -> LoadedDocument:
    text = Path(path).read_text(encoding="utf-8", errors="ignore")
    doc = LoadedDocument(source=path, doc_id=Path(path).stem)
    for i, para in enumerate(text.split("\n\n")):
        if para.strip():
            doc.elements.append(DocElement(type="text", content=para, section=f"para {i}"))
    return doc


def load_docx(path: str, ollama: OllamaClient, vision_model: str) -> LoadedDocument:
    from docx import Document

    doc_id = Path(path).stem
    doc = LoadedDocument(source=path, doc_id=doc_id)
    d = Document(path)
    for i, para in enumerate(d.paragraphs):
        if para.text.strip():
            doc.elements.append(DocElement(type="text", content=para.text, section=f"para {i}"))
    for ti, tbl in enumerate(d.tables):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in tbl.rows]
        doc.elements.append(DocElement(type="table", content="\n".join(rows), section=f"table {ti}"))
    return doc


def load_pptx(path: str, ollama: OllamaClient, vision_model: str) -> LoadedDocument:
    from pptx import Presentation

    doc_id = Path(path).stem
    doc = LoadedDocument(source=path, doc_id=doc_id)
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if p.text.strip():
                        texts.append(p.text)
        if texts:
            doc.elements.append(DocElement(type="text", content="\n".join(texts), page=i, section=f"slide {i}"))
        for shape in slide.shapes:
            if shape.shape_type == 13:
                try:
                    img = Image.open(io.BytesIO(shape.image.blob))
                    desc = describe_image_ollama(img, vision_model, ollama)
                    if desc:
                        doc.elements.append(DocElement(type="image", content=desc, page=i, section=f"slide {i}"))
                except Exception:
                    pass
    return doc


def load_xlsx(path: str) -> LoadedDocument:
    from openpyxl import load_workbook

    doc_id = Path(path).stem
    doc = LoadedDocument(source=path, doc_id=doc_id)
    wb = load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            if any(c is not None and str(c).strip() for c in row):
                rows.append(" | ".join("" if c is None else str(c) for c in row))
        if rows:
            doc.elements.append(DocElement(type="table", content="\n".join(rows), section=f"sheet {ws.title}"))
    wb.close()
    return doc


def load_document(path: str, ollama: OllamaClient, vision_model: str) -> LoadedDocument:
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        return load_pdf(str(p), ollama)
    if ext in {".md", ".markdown"}:
        return load_markdown(str(p))
    if ext in {".txt", ".log"}:
        return load_text(str(p))
    if ext == ".docx":
        return load_docx(str(p), ollama, vision_model)
    if ext == ".pptx":
        return load_pptx(str(p), ollama, vision_model)
    if ext == ".xlsx":
        return load_xlsx(str(p))
    raise ValueError(f"Unsupported file type: {ext}")


def load_paths(paths: Iterable[str], ollama: OllamaClient, vision_model: str) -> list[LoadedDocument]:
    docs: list[LoadedDocument] = []
    for p in paths:
        try:
            log.info(f"Loading {p}")
            docs.append(load_document(p, ollama, vision_model))
        except Exception as e:
            log.error(f"Failed to load {p}: {e}")
    return docs
